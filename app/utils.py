import shutil
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .lo_manager import LibreOfficeManager


CELLS = {
    "Combustão estacionária": ["E282", "E284"],
    "Combustão móvel": ["F805", "F807"],
    "Emissões fugitivas": ["E277", "F279"],
}

lo_manager = LibreOfficeManager()


def validar_categoria(categoria: str) -> None:
    if categoria not in CELLS:
        raise ValueError(
            f"Categoria inválida: {categoria}. "
            f"Categorias disponíveis: {list(CELLS)}"
        )


def calcular_emissao(
    arquivo_modelo: Path,
    arquivo_resultado: Path,
    changes: dict[str, dict[str, Any]],
) -> None:
    if not arquivo_modelo.exists():
        raise FileNotFoundError(
            f"Arquivo modelo não encontrado: {arquivo_modelo.resolve()}"
        )

    shutil.copy2(arquivo_modelo, arquivo_resultado)
    workbook = load_workbook(arquivo_resultado)

    try:
        for nome_aba, celulas in changes.items():
            if nome_aba not in workbook.sheetnames:
                raise ValueError(
                    f"A aba '{nome_aba}' não existe. "
                    f"Abas disponíveis: {workbook.sheetnames}"
                )

            worksheet = workbook[nome_aba]
            for celula, valor in celulas.items():
                worksheet[celula] = valor

        workbook.save(arquivo_resultado)
    finally:
        workbook.close()


def get_result_cells(
    input_path: Path,
    categoria: str,
) -> dict[str, Any]:
    validar_categoria(categoria)

    with tempfile.TemporaryDirectory() as temp_dir:
        recalculated_path = Path(temp_dir) / "resultado.xlsx"
        lo_manager.recalculate(input_path, recalculated_path)

        workbook = load_workbook(recalculated_path, data_only=True, read_only=True)
        try:
            if categoria not in workbook.sheetnames:
                raise ValueError(f"A aba '{categoria}' não existe na planilha.")

            sheet = workbook[categoria]
            return {cell: sheet[cell].value for cell in CELLS[categoria]}
        finally:
            workbook.close()
